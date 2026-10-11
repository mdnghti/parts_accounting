from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models.document import Document
from models.part import Part


# Стилі для Excel
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F497D")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
TOTAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
TOTAL_FONT = Font(name="Calibri", size=11, bold=True)
LOW_STOCK_FILL = PatternFill(start_color="FFD9D9", end_color="FFD9D9", fill_type="solid")
LOW_STOCK_FONT = Font(name="Calibri", size=11, color="9C0006", bold=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
DOUBLE_BOTTOM_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="double", color="000000"),
)


def _autofit_columns(ws, min_width: int = 10, max_width: int = 50) -> None:
    """Автоматичне підлаштування ширини колонок за вмістом."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_width), max_width)


def export_stock_to_excel(parts: List[Part], output_path: Path | str) -> str:
    """Звіт: Залишки на складі."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Залишки на складі"
    ws.views.sheetView[0].showGridLines = True

    # Заголовок звіту
    ws["A1"] = "ЗВІТ: ЗАЛИШКИ АВТОЗАПЧАСТИН НА СКЛАДІ"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Дата формування: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}"
    ws["A2"].font = SUBTITLE_FONT

    headers = [
        "№ з/п",
        "Артикул",
        "Назва запчастини",
        "Категорія",
        "Ціна (грн)",
        "Залишок",
        "Мін. залишок",
        "Од. вим.",
        "Сума (грн)",
        "Статус",
    ]

    header_row = 4
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_qty = 0
    total_amount = 0.0

    current_row = header_row + 1
    for idx, part in enumerate(parts, 1):
        row_amount = round(part.price * part.stock_quantity, 2)
        total_qty += part.stock_quantity
        total_amount += row_amount

        status = "ДЕФІЦИТ" if part.is_low_stock else "В нормі"

        ws.cell(row=current_row, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=2, value=part.part_number).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=3, value=part.name)
        ws.cell(row=current_row, column=4, value=part.category_name or "—")

        cell_price = ws.cell(row=current_row, column=5, value=part.price)
        cell_price.number_format = "#,##0.00"

        cell_stock = ws.cell(row=current_row, column=6, value=part.stock_quantity)
        cell_stock.alignment = Alignment(horizontal="center")

        cell_min = ws.cell(row=current_row, column=7, value=part.min_stock_quantity)
        cell_min.alignment = Alignment(horizontal="center")

        ws.cell(row=current_row, column=8, value=part.unit).alignment = Alignment(horizontal="center")

        cell_amount = ws.cell(row=current_row, column=9, value=row_amount)
        cell_amount.number_format = "#,##0.00"

        cell_status = ws.cell(row=current_row, column=10, value=status)
        cell_status.alignment = Alignment(horizontal="center")

        # Підсвічування рядка дефіциту
        if part.is_low_stock:
            cell_stock.fill = LOW_STOCK_FILL
            cell_stock.font = LOW_STOCK_FONT
            cell_status.fill = LOW_STOCK_FILL
            cell_status.font = LOW_STOCK_FONT

        for c in range(1, 11):
            ws.cell(row=current_row, column=c).border = THIN_BORDER

        current_row += 1

    # Підсумковий рядок
    ws.cell(row=current_row, column=1, value="")
    ws.cell(row=current_row, column=2, value="")
    ws.cell(row=current_row, column=3, value="РАЗОМ ПО СКЛАДУ:")
    ws.cell(row=current_row, column=3).font = TOTAL_FONT

    ws.cell(row=current_row, column=6, value=total_qty).alignment = Alignment(horizontal="center")
    ws.cell(row=current_row, column=6).font = TOTAL_FONT

    cell_tot_amt = ws.cell(row=current_row, column=9, value=round(total_amount, 2))
    cell_tot_amt.number_format = "#,##0.00"
    cell_tot_amt.font = TOTAL_FONT

    for c in range(1, 11):
        cell = ws.cell(row=current_row, column=c)
        cell.fill = TOTAL_FILL
        cell.border = DOUBLE_BOTTOM_BORDER

    _autofit_columns(ws)
    target = str(Path(output_path).resolve())
    wb.save(target)
    return target


def export_movement_to_excel(documents: List[Document], output_path: Path | str) -> str:
    """Звіт: Рух товарів за період (журнал накладних з деталізацією)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Рух товарів"
    ws.views.sheetView[0].showGridLines = True

    ws["A1"] = "ЗВІТ: РУХ ТОВАРІВ ЗА ПЕРІОД"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Дата формування: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}"
    ws["A2"].font = SUBTITLE_FONT

    headers = [
        "№ документу",
        "Тип операції",
        "Дата і час",
        "Контрагент",
        "Артикул",
        "Найменування запчастини",
        "Кількість",
        "Ціна (грн)",
        "Сума (грн)",
    ]

    header_row = 4
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    current_row = header_row + 1
    total_qty = 0
    total_amount = 0.0

    for doc in documents:
        for item in doc.items:
            total_qty += item.quantity
            total_amount += item.total_price

            ws.cell(row=current_row, column=1, value=doc.doc_number).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=2, value=doc.doc_type_ukr).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=3, value=doc.doc_date).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=4, value=doc.counterparty_name)
            ws.cell(row=current_row, column=5, value=item.part_number).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=6, value=item.part_name)

            ws.cell(row=current_row, column=7, value=item.quantity).alignment = Alignment(horizontal="center")

            cp = ws.cell(row=current_row, column=8, value=item.price)
            cp.number_format = "#,##0.00"

            ca = ws.cell(row=current_row, column=9, value=item.total_price)
            ca.number_format = "#,##0.00"

            for c in range(1, 10):
                ws.cell(row=current_row, column=c).border = THIN_BORDER

            current_row += 1

    # Підсумок
    ws.cell(row=current_row, column=1, value="")
    ws.cell(row=current_row, column=2, value="")
    ws.cell(row=current_row, column=3, value="")
    ws.cell(row=current_row, column=4, value="")
    ws.cell(row=current_row, column=6, value="ВСЬОГО:")
    ws.cell(row=current_row, column=6).font = TOTAL_FONT

    ws.cell(row=current_row, column=7, value=total_qty).alignment = Alignment(horizontal="center")
    ws.cell(row=current_row, column=7).font = TOTAL_FONT

    cta = ws.cell(row=current_row, column=9, value=round(total_amount, 2))
    cta.number_format = "#,##0.00"
    cta.font = TOTAL_FONT

    for c in range(1, 10):
        cell = ws.cell(row=current_row, column=c)
        cell.fill = TOTAL_FILL
        cell.border = DOUBLE_BOTTOM_BORDER

    _autofit_columns(ws)
    target = str(Path(output_path).resolve())
    wb.save(target)
    return target


def export_counterparties_to_excel(
    suppliers_data: List[Dict[str, Any]],
    customers_data: List[Dict[str, Any]],
    output_path: Path | str,
) -> str:
    """Звіт: Обороти по контрагентах (постачальники та покупці)."""
    wb = openpyxl.Workbook()

    # Лист 1: Постачальники
    ws1 = wb.active
    ws1.title = "Постачальники"
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "ЗВІТ: ПРИХОДИ ВІД ПОСТАЧАЛЬНИКІВ"
    ws1["A1"].font = TITLE_FONT
    ws1["A2"] = f"Дата формування: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}"
    ws1["A2"].font = SUBTITLE_FONT

    headers1 = ["№", "Постачальник", "Телефон", "Кількість поставок", "Отримано одиниць", "Сума поставок (грн)"]
    for col_idx, h in enumerate(headers1, 1):
        c = ws1.cell(row=4, column=col_idx, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")

    row = 5
    for idx, s in enumerate(suppliers_data, 1):
        ws1.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws1.cell(row=row, column=2, value=s.get("name"))
        ws1.cell(row=row, column=3, value=s.get("phone") or "—")
        ws1.cell(row=row, column=4, value=s.get("doc_count", 0)).alignment = Alignment(horizontal="center")
        ws1.cell(row=row, column=5, value=s.get("total_quantity", 0)).alignment = Alignment(horizontal="center")
        ca = ws1.cell(row=row, column=6, value=s.get("total_amount", 0.0))
        ca.number_format = "#,##0.00"
        for col_idx in range(1, 7):
            ws1.cell(row=row, column=col_idx).border = THIN_BORDER
        row += 1
    _autofit_columns(ws1)

    # Лист 2: Покупці
    ws2 = wb.create_sheet(title="Покупці")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "ЗВІТ: ПРОДАЖІ КЛІЄНТАМ"
    ws2["A1"].font = TITLE_FONT
    ws2["A2"] = f"Дата формування: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}"
    ws2["A2"].font = SUBTITLE_FONT

    headers2 = ["№", "Клієнт (покупець)", "Телефон", "Кількість замовлень", "Куплено одиниць", "Сума продажів (грн)"]
    for col_idx, h in enumerate(headers2, 1):
        c = ws2.cell(row=4, column=col_idx, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")

    row = 5
    for idx, cust in enumerate(customers_data, 1):
        ws2.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws2.cell(row=row, column=2, value=cust.get("name"))
        ws2.cell(row=row, column=3, value=cust.get("phone") or "—")
        ws2.cell(row=row, column=4, value=cust.get("doc_count", 0)).alignment = Alignment(horizontal="center")
        ws2.cell(row=row, column=5, value=cust.get("total_quantity", 0)).alignment = Alignment(horizontal="center")
        ca = ws2.cell(row=row, column=6, value=cust.get("total_amount", 0.0))
        ca.number_format = "#,##0.00"
        for col_idx in range(1, 7):
            ws2.cell(row=row, column=col_idx).border = THIN_BORDER
        row += 1
    _autofit_columns(ws2)

    target = str(Path(output_path).resolve())
    wb.save(target)
    return target


def export_invoice_to_excel(document: Document, output_path: Path | str) -> str:
    """Друкована накладна в Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Накладна"
    ws.views.sheetView[0].showGridLines = True

    doc_title = "ПРИБУТКОВА НАКЛАДНА" if document.doc_type == "incoming" else "ВИДАТКОВА НАКЛАДНА"
    ws["A1"] = f"{doc_title} № {document.doc_number}"
    ws["A1"].font = TITLE_FONT

    ws["A2"] = f"від {document.doc_date}"
    ws["A2"].font = Font(name="Calibri", size=12, bold=True)

    ws["A4"] = "Одержувач / Постачальник:"
    ws["B4"] = document.counterparty_name
    ws["A4"].font = Font(bold=True)

    ws["A5"] = "Відповідальна особа:"
    ws["B5"] = document.user_name
    ws["A5"].font = Font(bold=True)

    headers = ["№", "Артикул", "Найменування товару", "Од.", "Кількість", "Ціна (грн)", "Сума (грн)"]
    header_row = 7
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col_idx, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")

    row = header_row + 1
    for idx, item in enumerate(document.items, 1):
        ws.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2, value=item.part_number).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=3, value=item.part_name)
        ws.cell(row=row, column=4, value=item.unit).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=5, value=item.quantity).alignment = Alignment(horizontal="center")

        cp = ws.cell(row=row, column=6, value=item.price)
        cp.number_format = "#,##0.00"

        ca = ws.cell(row=row, column=7, value=item.total_price)
        ca.number_format = "#,##0.00"

        for c in range(1, 8):
            ws.cell(row=row, column=c).border = THIN_BORDER
        row += 1

    # Всього
    ws.cell(row=row, column=3, value="ВСЬОГО ДО СПЛАТИ:")
    ws.cell(row=row, column=3).font = TOTAL_FONT

    ws.cell(row=row, column=5, value=document.total_quantity).alignment = Alignment(horizontal="center")
    ws.cell(row=row, column=5).font = TOTAL_FONT

    tot_c = ws.cell(row=row, column=7, value=document.total_amount)
    tot_c.number_format = "#,##0.00"
    tot_c.font = TOTAL_FONT

    for c in range(1, 8):
        cell = ws.cell(row=row, column=c)
        cell.fill = TOTAL_FILL
        cell.border = DOUBLE_BOTTOM_BORDER

    # Підписи
    row += 3
    ws.cell(row=row, column=2, value="Здав: ___________________")
    ws.cell(row=row, column=5, value="Прийняв: ___________________")

    _autofit_columns(ws)
    target = str(Path(output_path).resolve())
    wb.save(target)
    return target
