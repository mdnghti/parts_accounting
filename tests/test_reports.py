from pathlib import Path
import openpyxl
import pytest

from models import Document, DocumentItem, Part
from reports.excel_generator import (
    export_counterparties_to_excel,
    export_invoice_to_excel,
    export_movement_to_excel,
    export_stock_to_excel,
)
from reports.pdf_generator import (
    export_counterparties_to_pdf,
    export_invoice_to_pdf,
    export_movement_to_pdf,
    export_stock_to_pdf,
)
from reports.services import generate_report


@pytest.fixture
def sample_parts():
    return [
        Part(
            id=1,
            part_number="OF-1001",
            name="Фільтр масляний Mann W67/2",
            category_id=1,
            price=280.0,
            stock_quantity=15,
            min_stock_quantity=5,
            unit="шт.",
            category_name="Фільтри",
        ),
        Part(
            id=2,
            part_number="BP-5001",
            name="Колодки гальмівні передні",
            category_id=2,
            price=1200.0,
            stock_quantity=2,  # Дефіцит!
            min_stock_quantity=4,
            unit="компл.",
            category_name="Гальмівна система",
        ),
    ]


@pytest.fixture
def sample_document():
    return Document(
        id=1,
        doc_type="incoming",
        doc_number="ПР-2026-0001",
        doc_date="2026-10-11 12:00:00",
        counterparty_id=1,
        user_id=1,
        notes="Тестова поставка",
        counterparty_name="ТОВ 'АвтоТрейд Постач'",
        user_name="Адміністратор",
        items=[
            DocumentItem(id=1, document_id=1, part_id=1, quantity=10, price=250.0, part_number="OF-1001", part_name="Фільтр масляний", unit="шт."),
            DocumentItem(id=2, document_id=1, part_id=2, quantity=4, price=1000.0, part_number="BP-5001", part_name="Колодки гальмівні", unit="компл."),
        ],
    )


def test_excel_stock_export(sample_parts, tmp_path: Path):
    out_file = tmp_path / "stock_test.xlsx"
    res_path = export_stock_to_excel(sample_parts, out_file)
    assert Path(res_path).exists()
    assert Path(res_path).stat().st_size > 0

    # Читання назад через openpyxl для перевірки структури
    wb = openpyxl.load_workbook(res_path)
    assert "Залишки на складі" in wb.sheetnames
    ws = wb["Залишки на складі"]
    assert "ЗВІТ: ЗАЛИШКИ АВТОЗАПЧАСТИН НА СКЛАДІ" in ws["A1"].value
    # Рядок заголовків
    assert ws["B4"].value == "Артикул"
    # Рядок 1 запчастини
    assert ws["B5"].value == "OF-1001"
    # Рядок 2 запчастини (дефіцит)
    assert ws["B6"].value == "BP-5001"
    assert ws["J6"].value == "ДЕФІЦИТ"


def test_excel_movement_export(sample_document, tmp_path: Path):
    out_file = tmp_path / "movement_test.xlsx"
    res_path = export_movement_to_excel([sample_document], out_file)
    assert Path(res_path).exists()
    assert Path(res_path).stat().st_size > 0

    wb = openpyxl.load_workbook(res_path)
    assert "Рух товарів" in wb.sheetnames


def test_excel_counterparties_export(tmp_path: Path):
    out_file = tmp_path / "counterparties_test.xlsx"
    sups = [{"name": "Постачальник 1", "phone": "123", "doc_count": 2, "total_quantity": 20, "total_amount": 5000.0}]
    custs = [{"name": "Клієнт 1", "phone": "456", "doc_count": 1, "total_quantity": 5, "total_amount": 2500.0}]

    res_path = export_counterparties_to_excel(sups, custs, out_file)
    assert Path(res_path).exists()

    wb = openpyxl.load_workbook(res_path)
    assert "Постачальники" in wb.sheetnames
    assert "Покупці" in wb.sheetnames


def test_excel_invoice_export(sample_document, tmp_path: Path):
    out_file = tmp_path / "invoice_test.xlsx"
    res_path = export_invoice_to_excel(sample_document, out_file)
    assert Path(res_path).exists()

    wb = openpyxl.load_workbook(res_path)
    assert "Накладна" in wb.sheetnames
    ws = wb["Накладна"]
    assert sample_document.doc_number in ws["A1"].value


def test_pdf_exports(sample_parts, sample_document, tmp_path: Path):
    # 1. Залишки в PDF
    pdf_stock = tmp_path / "stock.pdf"
    res1 = export_stock_to_pdf(sample_parts, pdf_stock)
    assert Path(res1).exists()
    assert Path(res1).stat().st_size > 1000
    with open(res1, "rb") as f:
        assert f.read(5) == b"%PDF-"

    # 2. Рух в PDF
    pdf_move = tmp_path / "movement.pdf"
    res2 = export_movement_to_pdf([sample_document], pdf_move)
    assert Path(res2).exists()
    assert Path(res2).stat().st_size > 1000

    # 3. Контрагенти в PDF
    pdf_cp = tmp_path / "counterparties.pdf"
    sups = [{"name": "Постачальник 1", "phone": "123", "doc_count": 1, "total_quantity": 10, "total_amount": 2000.0}]
    custs = [{"name": "Клієнт 1", "phone": "456", "doc_count": 1, "total_quantity": 2, "total_amount": 500.0}]
    res3 = export_counterparties_to_pdf(sups, custs, pdf_cp)
    assert Path(res3).exists()

    # 4. Друкована накладна в PDF
    pdf_inv = tmp_path / "invoice.pdf"
    res4 = export_invoice_to_pdf(sample_document, pdf_inv)
    assert Path(res4).exists()
    assert Path(res4).stat().st_size > 1000


def test_report_services_dispatcher(sample_document, tmp_path: Path):
    out_xlsx = tmp_path / "disp_stock.xlsx"
    r1 = generate_report("stock", "excel", out_xlsx)
    assert Path(r1).exists()

    out_pdf = tmp_path / "disp_invoice.pdf"
    r2 = generate_report("invoice", "pdf", out_pdf, document=sample_document)
    assert Path(r2).exists()
